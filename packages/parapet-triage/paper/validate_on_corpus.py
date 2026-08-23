"""Score the shipped pattern set against the published labelled corpus.

The performance pattern set was derived and validated on a corpus of
hand-tagged sentences and issue reports. That corpus is published, so the
tool's accuracy on it is a number the tool can state rather than claim.

    python paper/validate_on_corpus.py --workbook "Manual Tagging.xlsx" \
        --out ../../docs/accuracy.md

Run it after changing a pattern definition, a threshold or the backend.
The numbers it prints are what the README quotes; if they move, the README
moves with them.

The corpus is CC BY 4.0 (doi:10.5281/zenodo.10944186) and is not
redistributed with the package. Download it to run this script.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from parapet_triage import Triager, load_pattern_set   # noqa: E402
from parapet_triage.dsl import compile_set             # noqa: E402

SENTENCE_SHEETS = ("Dataset-1 Sen", "Dataset-2 Sen", "Dataset-3 Sen")
ISSUE_SHEETS = ("Dataset-1 Issue", "Dataset-2 Issue", "Dataset-3 Issue")


def to_binary(value) -> int:
    text = str(value).strip().lower() if value is not None else ""
    if text in ("yes", "y", "1", "true"):
        return 1
    if text in ("no", "n", "0", "false"):
        return 0
    return -1


def load_sentences(workbook) -> Dict[str, List[Tuple[str, int]]]:
    """``key, no, text, label`` sheets, plus Dataset-3's side-by-side layout."""
    out: Dict[str, List[Tuple[str, int]]] = {}
    for sheet in SENTENCE_SHEETS:
        if sheet not in workbook.sheetnames:
            continue
        rows = [list(r) for r in workbook[sheet].iter_rows(values_only=True)]
        items: List[Tuple[str, int]] = []
        if sheet.endswith("3 Sen"):
            # Two sources side by side: Bugzilla at column 0, Redmine at 7.
            for row in rows[2:]:
                for base in (0, 7):
                    if len(row) > base + 3 and row[base + 2]:
                        label = to_binary(row[base + 3])
                        if label >= 0:
                            items.append((str(row[base + 2]).strip(), label))
        else:
            for row in rows[1:]:
                if len(row) > 3 and row[2]:
                    label = to_binary(row[3])
                    if label >= 0:
                        items.append((str(row[2]).strip(), label))
        out[sheet] = items
    return out


def load_issues(workbook) -> Dict[str, List[Tuple[str, str, int]]]:
    """``key, text, Manual Label`` sheets."""
    out: Dict[str, List[Tuple[str, str, int]]] = {}
    for sheet in ISSUE_SHEETS:
        if sheet not in workbook.sheetnames:
            continue
        rows = [list(r) for r in workbook[sheet].iter_rows(values_only=True)]
        header_index = 0
        for index, row in enumerate(rows[:4]):
            joined = " ".join(str(cell).lower() for cell in row[:6] if cell)
            if "manual" in joined and ("label" in joined or "tag" in joined):
                header_index = index
                break
        header = [str(cell).strip().lower() if cell else ""
                  for cell in rows[header_index]]

        def column(*names, default=None):
            for name in names:
                if name in header:
                    return header.index(name)
            return default

        key_col = column("key", "id", "bug id", default=0)
        text_col = column("text", "summary", "description", default=1)
        label_col = column("manual label", "author-1 label", "author-1 tag",
                           default=2)
        items = []
        for row in rows[header_index + 1:]:
            if not row or len(row) <= max(key_col, text_col, label_col):
                continue
            label = to_binary(row[label_col])
            if label < 0 or not row[text_col]:
                continue
            items.append((str(row[key_col]).strip(), str(row[text_col]).strip(),
                          label))
        out[sheet] = items
    return out


def score(true_labels: Sequence[int], predicted: Sequence[int]) -> Dict[str, float]:
    tp = sum(1 for t, p in zip(true_labels, predicted) if t and p)
    fp = sum(1 for t, p in zip(true_labels, predicted) if not t and p)
    fn = sum(1 for t, p in zip(true_labels, predicted) if t and not p)
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    return {"n": len(true_labels), "positives": sum(true_labels),
            "tp": tp, "fp": fp, "fn": fn,
            "precision": precision, "recall": recall, "f1": f1}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--patterns", default="performance")
    parser.add_argument("--backend", default="simple", choices=("simple", "spacy"))
    parser.add_argument("--sentence-threshold", type=float, default=None)
    parser.add_argument("--issue-threshold", type=float, default=None)
    parser.add_argument("--sweep", action="store_true",
                        help="sweep thresholds instead of scoring one setting")
    parser.add_argument("--limit", type=int, default=0,
                        help="cap sentences per sheet, for a quick check")
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args()

    import openpyxl
    workbook = openpyxl.load_workbook(args.workbook, read_only=True,
                                      data_only=True)
    sentences = load_sentences(workbook)
    issues = load_issues(workbook)
    workbook.close()

    pattern_set = load_pattern_set(args.patterns)
    _compiled, compile_report = compile_set(pattern_set.patterns)
    print(f"pattern set {pattern_set.set_id}: {len(pattern_set)} patterns, "
          f"{compile_report['n_repaired']} repaired, "
          f"{compile_report['n_need_pos']} need a tagger")

    kwargs = {}
    if args.sentence_threshold is not None:
        kwargs["sentence_threshold"] = args.sentence_threshold
    if args.issue_threshold is not None:
        kwargs["issue_threshold"] = args.issue_threshold
    triager = Triager(pattern_set=pattern_set, backend=args.backend, **kwargs)

    lines: List[str] = []

    if args.sweep:
        _sweep(triager, sentences, args.limit)
        return

    print("\n=== sentence level")
    for sheet, items in sentences.items():
        rows = items[: args.limit] if args.limit else items
        if not rows:
            continue
        annotated = [triager.backend.annotate(text) for text, _ in rows]
        predicted = [1 if triager.score_sentence(s).is_match else 0
                     for s in annotated]
        result = score([label for _, label in rows], predicted)
        lines.append(_row(sheet, result))
        print(_row(sheet, result))

    print("\n=== issue level")
    for sheet, items in issues.items():
        rows = items[: args.limit] if args.limit else items
        if not rows:
            continue
        predicted = [1 if triager.classify_text(text, key).is_match else 0
                     for key, text, _ in rows]
        result = score([label for _, _, label in rows], predicted)
        lines.append(_row(sheet, result))
        print(_row(sheet, result))

    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(
            "# Accuracy on the published corpus\n\n"
            f"Pattern set `{pattern_set.set_id}`, `{args.backend}` backend, "
            f"sentence threshold {triager.sentence_threshold}, "
            f"issue threshold {triager.issue_threshold}.\n\n"
            "| sheet | n | positives | precision | recall | F1 |\n"
            "| --- | ---: | ---: | ---: | ---: | ---: |\n"
            + "\n".join(lines) + "\n\n"
            "Reproduce with `python paper/validate_on_corpus.py --workbook "
            "\"Manual Tagging.xlsx\"`. The corpus is CC BY 4.0, "
            "doi:10.5281/zenodo.10944186, and is not shipped with the "
            "package.\n",
            encoding="utf-8")
        print(f"\nwrote {args.out}")


def _row(sheet: str, result: Dict[str, float]) -> str:
    return (f"| {sheet} | {result['n']} | {result['positives']} | "
            f"{result['precision']:.2f} | {result['recall']:.2f} | "
            f"{result['f1']:.2f} |")


def _sweep(triager, sentences, limit: int) -> None:
    """Print F1 across sentence thresholds, so a default can be justified."""
    print("\nsentence threshold sweep")
    print("| threshold | " + " | ".join(sentences) + " |")
    cache = {}
    for sheet, items in sentences.items():
        rows = items[:limit] if limit else items
        cache[sheet] = [(triager.score_sentence(triager.backend.annotate(text)).score,
                         label) for text, label in rows]
    for threshold in (0.5, 0.7, 0.85, 0.9, 1.0, 1.2, 1.5, 1.8, 2.0, 2.5):
        cells = []
        for sheet in sentences:
            scored = cache[sheet]
            predicted = [1 if value >= threshold else 0 for value, _ in scored]
            result = score([label for _, label in scored], predicted)
            cells.append(f"{result['f1']:.2f}")
        print(f"| {threshold:<9} | " + " | ".join(cells) + " |")


if __name__ == "__main__":
    main()
